from __future__ import annotations

import mimetypes
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi import UploadFile

from .database import Database
from .jobs import JobService
from .storage import LocalStorage, StorageError


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


_ALLOWED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".png", ".jpg", ".jpeg", ".webp"}


class DocumentService:
    def __init__(
        self,
        db: Database,
        storage: LocalStorage,
        jobs: JobService,
        *,
        max_upload_bytes: int,
        max_document_pages: int = 4000,
        max_extracted_chars: int = 20_000_000,
    ):
        self.db = db
        self.storage = storage
        self.jobs = jobs
        self.max_upload_bytes = max_upload_bytes
        self.max_document_pages = max_document_pages
        self.max_extracted_chars = max_extracted_chars

    async def upload(
        self,
        *,
        user_id: str,
        upload: UploadFile,
        course_id: str | None = None,
    ) -> dict[str, Any]:
        filename = Path(upload.filename or "upload.bin").name
        ext = Path(filename).suffix.lower()
        if ext not in _ALLOWED_EXTENSIONS:
            raise ValueError("Unsupported file type")
        content_type = upload.content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
        stored = await self.storage.save_upload(upload, max_bytes=self.max_upload_bytes)
        try:
            self._validate_signature(self.storage.path_for(stored.storage_key), ext)
        except Exception:
            self.storage.delete(stored.storage_key)
            raise
        duplicate = self.db.fetchone(
            "SELECT * FROM documents WHERE owner_user_id=? AND coalesce(course_id,'')=coalesce(?,'') AND sha256=? ORDER BY version DESC LIMIT 1",
            (user_id, course_id, stored.sha256),
        )
        if duplicate:
            self.storage.delete(stored.storage_key)
            result = self._row(duplicate)
            result["duplicate"] = True
            return result
        document_id = str(uuid4())
        now = _now()
        self.db.execute(
            """
            INSERT INTO documents(document_id,owner_user_id,course_id,filename,content_type,size_bytes,sha256,storage_key,version,status,created_at,updated_at)
            VALUES(?,?,?,?,?,?,?,?,1,'uploaded',?,?)
            """,
            (
                document_id,
                user_id,
                course_id,
                filename,
                content_type,
                stored.size_bytes,
                stored.sha256,
                stored.storage_key,
                now,
                now,
            ),
        )
        job = self.jobs.create(
            user_id=user_id,
            kind="document_inspect",
            resource_id=document_id,
            priority=2,
            idempotency_key=f"inspect:{document_id}:v1",
        )
        # Native extraction is deliberately synchronous only for the local release
        # adapter. The durable job record preserves the contract for external workers.
        self.jobs.update(user_id, job["job_id"], status="running", progress=0.1, increment_attempt=True)
        try:
            result = self._native_extract(user_id, document_id)
            self.jobs.update(user_id, job["job_id"], status="succeeded", progress=1.0, result=result)
        except Exception as exc:
            self.db.execute(
                "UPDATE documents SET status='failed_retryable',error_message=?,updated_at=? WHERE document_id=? AND owner_user_id=?",
                (str(exc)[:1000], _now(), document_id, user_id),
            )
            self.jobs.update(user_id, job["job_id"], status="failed", error=str(exc))
        return self.get(user_id, document_id)

    def get(self, user_id: str, document_id: str) -> dict[str, Any]:
        row = self.db.fetchone("SELECT * FROM documents WHERE document_id=? AND owner_user_id=?", (document_id, user_id))
        if not row:
            raise ValueError("Unknown document")
        return self._row(row)

    def list(self, user_id: str, course_id: str | None = None, limit: int = 100) -> list[dict[str, Any]]:
        if course_id:
            rows = self.db.fetchall(
                "SELECT * FROM documents WHERE owner_user_id=? AND course_id=? ORDER BY updated_at DESC LIMIT ?",
                (user_id, course_id, max(1, min(200, limit))),
            )
        else:
            rows = self.db.fetchall(
                "SELECT * FROM documents WHERE owner_user_id=? ORDER BY updated_at DESC LIMIT ?",
                (user_id, max(1, min(200, limit))),
            )
        return [self._row(r) for r in rows]

    def delete(self, user_id: str, document_id: str) -> None:
        row = self.db.fetchone(
            "SELECT storage_key FROM documents WHERE document_id=? AND owner_user_id=?",
            (document_id, user_id),
        )
        if not row:
            raise ValueError("Unknown document")
        # Delete derived retrieval content in the same DB transaction so a deleted
        # private document cannot remain searchable after its metadata disappears.
        with self.db.transaction() as conn:
            conn.execute(
                "DELETE FROM knowledge_chunks WHERE document_id=? AND owner_user_id=?",
                (document_id, user_id),
            )
            conn.execute(
                "DELETE FROM jobs WHERE resource_id=? AND owner_user_id=?",
                (document_id, user_id),
            )
            conn.execute(
                "DELETE FROM documents WHERE document_id=? AND owner_user_id=?",
                (document_id, user_id),
            )
        self.storage.delete(row["storage_key"])

    def page(self, user_id: str, document_id: str, page_no: int) -> dict[str, Any]:
        self.get(user_id, document_id)
        row = self.db.fetchone("SELECT * FROM document_pages WHERE document_id=? AND page_no=?", (document_id, page_no))
        if not row:
            raise ValueError("Page is not available")
        return {
            "document_id": document_id,
            "page_no": page_no,
            "status": row["status"],
            "text": row["extracted_text"],
            "provenance": self.db.loads(row["provenance_json"], {}),
        }

    def _native_extract(self, user_id: str, document_id: str) -> dict[str, Any]:
        row = self.db.fetchone(
            "SELECT * FROM documents WHERE document_id=? AND owner_user_id=?",
            (document_id, user_id),
        )
        if not row:
            raise ValueError("Unknown document")
        doc = dict(row)
        path = self.storage.path_for(doc["storage_key"])
        ext = Path(doc["filename"]).suffix.lower()
        pages: list[str] = []
        provenance = "native"
        if ext in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            if len(text) > self.max_extracted_chars:
                raise ValueError("Document exceeds configured extracted-text safety limit")
            pages = [text]
        elif ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            if reader.is_encrypted:
                self.db.execute(
                    "UPDATE documents SET status='password_required',updated_at=? WHERE document_id=?",
                    (_now(), document_id),
                )
                return {"status": "password_required", "pages": 0}
            if len(reader.pages) > self.max_document_pages:
                raise ValueError("Document exceeds configured page safety limit")
            total_chars = 0
            for page in reader.pages:
                text = page.extract_text() or ""
                total_chars += len(text)
                if total_chars > self.max_extracted_chars:
                    raise ValueError("Document exceeds configured extracted-text safety limit")
                pages.append(text)
        elif ext == ".docx":
            from docx import Document

            d = Document(str(path))
            text = "\n".join(p.text for p in d.paragraphs)
            if len(text) > self.max_extracted_chars:
                raise ValueError("Document exceeds configured extracted-text safety limit")
            pages = [text]
        elif ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(path))
            if len(prs.slides) > self.max_document_pages:
                raise ValueError("Document exceeds configured page safety limit")
            total_chars = 0
            for slide in prs.slides:
                texts: list[str] = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        texts.append(shape.text)
                text = "\n".join(texts)
                total_chars += len(text)
                if total_chars > self.max_extracted_chars:
                    raise ValueError("Document exceeds configured extracted-text safety limit")
                pages.append(text)
        else:
            # Image OCR remains a separate on-demand service; upload is usable even
            # before OCR has completed.
            provenance = "ocr_pending"
            pages = [""]

        now = _now()
        with self.db.transaction() as conn:
            conn.execute("DELETE FROM document_pages WHERE document_id=?", (document_id,))
            conn.execute("DELETE FROM knowledge_chunks WHERE document_id=? AND owner_user_id=?", (document_id, user_id))
            course = conn.execute("SELECT subject_id FROM courses WHERE course_id=? AND owner_user_id=?", (doc["course_id"], user_id)).fetchone() if doc["course_id"] else None
            subject_id = course["subject_id"] if course else None
            for i, text in enumerate(pages, start=1):
                status = "native_ready" if text.strip() else "ocr_pending"
                conn.execute(
                    "INSERT INTO document_pages(document_id,page_no,status,extracted_text,provenance_json,updated_at) VALUES(?,?,?,?,?,?)",
                    (document_id, i, status, text, self.db.dumps({"method": provenance}), now),
                )
                if subject_id and text.strip():
                    # Page-sized source chunks are deliberately simple in v0.01;
                    # future structural chunking can replace them without changing
                    # ownership/provenance semantics.
                    conn.execute(
                        """INSERT INTO knowledge_chunks(
                            chunk_id,owner_user_id,course_id,document_id,source_id,subject_id,unit,topic,concept_ids_json,prerequisite_ids_json,difficulty,content_type,content,citation_location,verified,version,created_at
                        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            f"doc:{document_id}:p:{i}", user_id, doc["course_id"], document_id, document_id, subject_id,
                            doc["filename"], None, "[]", "[]", 0.5, "uploaded_material", text[:100000],
                            f"{doc['filename']} · page {i}", 0, str(doc["version"]), now,
                        ),
                    )
            ready_partial = any(text.strip() for text in pages)
            if ready_partial and all(text.strip() for text in pages):
                status = "ready"
            elif ready_partial:
                status = "ready_partial"
            else:
                status = "processing"
            conn.execute(
                "UPDATE documents SET status=?,error_message=NULL,updated_at=? WHERE document_id=?",
                (status, now, document_id),
            )
        pending_pages = [i for i, text in enumerate(pages, start=1) if not text.strip()]
        if pending_pages:
            self.jobs.create(
                user_id=user_id,
                kind="ocr_page_range",
                resource_id=document_id,
                priority=2,
                payload={"document_id": document_id, "page_numbers": pending_pages[: self.max_document_pages]},
                idempotency_key=f"ocr:{document_id}:v{doc['version']}",
            )
        return {"status": status, "pages": len(pages), "ocr_pending_pages": len(pending_pages)}

    @staticmethod
    def _validate_signature(path: Path, ext: str) -> None:
        # Extension is only a routing hint. Validate common container/image magic
        # before any parser sees attacker-controlled bytes. Text formats are the
        # deliberate exception because they have no stable magic prefix.
        if ext in {".txt", ".md"}:
            return
        with path.open("rb") as fh:
            head = fh.read(16)
        if ext == ".pdf" and not head.startswith(b"%PDF-"):
            raise ValueError("File content does not match PDF extension")
        if ext in {".docx", ".pptx"}:
            if not head.startswith(b"PK"):
                raise ValueError("File content does not match Office document extension")
            try:
                with zipfile.ZipFile(path) as archive:
                    entries = archive.infolist()
                    if len(entries) > 10_000:
                        raise ValueError("Office document contains too many archive entries")
                    total_uncompressed = 0
                    for entry in entries:
                        if entry.file_size > 128 * 1024 * 1024:
                            raise ValueError("Office document contains an oversized archive entry")
                        total_uncompressed += entry.file_size
                        if total_uncompressed > 512 * 1024 * 1024:
                            raise ValueError("Office document expands beyond the safety limit")
                        normalized = entry.filename.replace("\\", "/")
                        if normalized.startswith("/") or "../" in f"/{normalized}":
                            raise ValueError("Office document contains an unsafe archive path")
            except zipfile.BadZipFile as exc:
                raise ValueError("Office document container is invalid") from exc
        if ext == ".png" and not head.startswith(b"\x89PNG\r\n\x1a\n"):
            raise ValueError("File content does not match PNG extension")
        if ext in {".jpg", ".jpeg"} and not head.startswith(b"\xff\xd8\xff"):
            raise ValueError("File content does not match JPEG extension")
        if ext == ".webp" and not (head.startswith(b"RIFF") and head[8:12] == b"WEBP"):
            raise ValueError("File content does not match WEBP extension")

    def _row(self, row: Any) -> dict[str, Any]:
        return {
            "document_id": row["document_id"],
            "course_id": row["course_id"],
            "filename": row["filename"],
            "content_type": row["content_type"],
            "size_bytes": row["size_bytes"],
            "sha256": row["sha256"],
            "version": row["version"],
            "status": row["status"],
            "error": row["error_message"],
            "created_at": row["created_at"],
            "updated_at": row["updated_at"],
        }
